from __future__ import annotations

import codecs
import csv
import importlib
import json
import tempfile
from concurrent.futures import CancelledError
from pathlib import Path
from threading import Event
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.writer.excel import ExcelWriter
from openpyxl.worksheet import _writer as worksheet_writer


FIELDS = ("nick_name", "remark", "phone", "username", "alias", "description")
HEADERS = ["\u6635\u79f0", "\u5907\u6ce8", "\u624b\u673a\u53f7", "\u5fae\u4fe1 ID", "\u5fae\u4fe1\u53f7", "\u63cf\u8ff0"]
STEM = "\u5fae\u4fe1\u8054\u7cfb\u4eba"


def _export_module():
    try:
        return importlib.import_module("app.contacts.export")
    except ModuleNotFoundError as exc:
        pytest.fail(f"Contact export implementation is missing: {exc}")


def _records():
    return [dict(zip(FIELDS, ("\u5f20\u4e09\U0001f600", "\u4e00,\u4e8c\n\u4e09", "0013800000000", "00012345678901234567890", "0", "")), category="friend")]


def _raw(records):
    return [{key: row[key] for key in FIELDS} for row in records]


def test_all_formats_use_chinese_paths_and_identical_canonical_data(tmp_path):
    module = _export_module()
    target = tmp_path / "\u4e2d\u6587\u76ee\u5f55"
    records = _records()
    paths = module.export_contacts(records, "all", target)
    assert paths == [str(target / f"{STEM}.{fmt}") for fmt in ("csv", "json", "xlsx")]
    assert sorted(path.name for path in target.iterdir()) == sorted(f"{STEM}.{fmt}" for fmt in ("csv", "json", "xlsx"))
    csv_path, json_path, xlsx_path = map(Path, paths)
    assert csv_path.read_bytes().startswith(codecs.BOM_UTF8)
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows == [HEADERS, [records[0][key] for key in FIELDS]]
    assert json.loads(json_path.read_text(encoding="utf-8")) == _raw(records)
    assert list(json.loads(json_path.read_text(encoding="utf-8"))[0]) == list(FIELDS)
    workbook = load_workbook(xlsx_path)
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert [cell.value or "" for cell in sheet[2]] == [records[0][key] for key in FIELDS]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F2"
    assert all(cell.font.bold and cell.fill.fill_type == "solid" for cell in sheet[1])
    assert all(sheet.cell(2, column).number_format == "@" for column in (3, 4, 5))
    assert sheet.cell(2, 3).data_type == "s"
    workbook.close()


@pytest.mark.parametrize("value", ["=SUM(1,2)", "+00123", "-01", "@SUM(A1)", "\t=1", "\r=1", "\n=1", "  =1"])
def test_csv_prefixes_formula_injection_and_json_xlsx_remain_lossless(tmp_path, value):
    records = [dict.fromkeys(FIELDS, value)]
    paths = _export_module().export_contacts(records, "all", tmp_path)
    with Path(paths[0]).open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[1] == ["'" + value] * 6
    assert json.loads(Path(paths[1]).read_text(encoding="utf-8")) == records
    workbook = load_workbook(paths[2], data_only=False)
    assert [cell.value for cell in workbook.active[2]] == [value] * 6
    assert all(cell.data_type == "s" for cell in workbook.active[2])
    workbook.close()


@pytest.mark.parametrize("fmt", ("csv", "json", "xlsx", "all"))
def test_empty_contacts_fail_without_creating_outputs(tmp_path, fmt):
    module = _export_module()
    with pytest.raises(ValueError, match="(?i)empty|zero|no contacts|\u8054\u7cfb\u4eba"):
        module.export_contacts([], fmt, tmp_path / "new")
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("fmt", ("csv", "json", "xlsx"))
def test_individual_format_writes_exact_target_and_requires_overwrite(tmp_path, fmt):
    target = tmp_path / f"\u81ea\u5b9a\u4e49.{fmt}"
    target.write_bytes(b"original")
    module = _export_module()
    with pytest.raises(FileExistsError):
        module.export_contacts(_records(), fmt, target)
    assert target.read_bytes() == b"original"
    assert module.export_contacts(_records(), fmt, target, overwrite=True) == [str(target)]
    assert target.read_bytes() != b"original"
    assert list(tmp_path.iterdir()) == [target]


def test_all_formats_refuse_any_existing_destination_before_writing(tmp_path):
    target = tmp_path / f"{STEM}.json"
    target.write_bytes(b"original")
    unrelated = tmp_path / "keep.txt"
    unrelated.write_bytes(b"keep")
    with pytest.raises(FileExistsError):
        _export_module().export_contacts(_records(), "all", tmp_path)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {target.name: b"original", "keep.txt": b"keep"}


def test_progress_reports_bounded_batches_and_records_are_frozen(tmp_path):
    records = _records() * 300
    progress = []

    def report(done, total):
        progress.append((done, total))
        records[0]["nick_name"] = "mutated"

    paths = _export_module().export_contacts(records, "all", tmp_path, progress=report)
    assert progress[0] == (0, 900)
    assert progress[-1] == (900, 900)
    assert all(a[0] <= b[0] and b[0] - a[0] <= 256 for a, b in zip(progress, progress[1:]))
    assert all(row["nick_name"] == "\u5f20\u4e09\U0001f600" for row in json.loads(Path(paths[1]).read_text(encoding="utf-8")))


@pytest.mark.parametrize("fmt", ("csv", "json", "xlsx", "all"))
def test_pre_cancel_does_not_touch_existing_files(tmp_path, fmt):
    cancel = Event()
    cancel.set()
    original = tmp_path / (f"{STEM}.csv" if fmt == "all" else f"out.{fmt}")
    original.write_bytes(b"original")
    target = tmp_path if fmt == "all" else original
    with pytest.raises(CancelledError):
        _export_module().export_contacts(_records(), fmt, target, overwrite=True, cancel=cancel)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {original.name: b"original"}


@pytest.mark.parametrize("fmt", ("csv", "json", "xlsx", "all"))
def test_cancel_during_staging_cleans_temps_and_preserves_originals(tmp_path, fmt):
    module = _export_module()
    cancel = Event()
    original = tmp_path / (f"{STEM}.csv" if fmt == "all" else f"out.{fmt}")
    original.write_bytes(b"original")
    target = tmp_path if fmt == "all" else original

    def report(done, total):
        if done >= 128:
            cancel.set()

    with pytest.raises(CancelledError):
        module.export_contacts(_records() * 500, fmt, target, overwrite=True, cancel=cancel, progress=report)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {original.name: b"original"}


@pytest.mark.parametrize("existing", (False, True))
def test_partial_commit_failure_rolls_back_every_format(tmp_path, monkeypatch, existing):
    module = _export_module()
    before = {"keep.txt": b"unrelated"}
    if existing:
        before.update({f"{STEM}.{fmt}": f"old {fmt}".encode() for fmt in ("csv", "json", "xlsx")})
    for name, content in before.items():
        (tmp_path / name).write_bytes(content)
    real_replace = module.os.replace

    def replace(source, destination):
        if Path(destination).name == f"{STEM}.json" and Path(source).suffix == ".tmp":
            raise PermissionError("injected second commit failure")
        return real_replace(source, destination)

    monkeypatch.setattr(module.os, "replace", replace)
    with pytest.raises(PermissionError, match="second commit"):
        module.export_contacts(_records(), "all", tmp_path, overwrite=existing)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


def test_cancel_after_first_commit_rolls_back_replaced_originals(tmp_path, monkeypatch):
    module = _export_module()
    cancel = Event()
    before = {f"{STEM}.{fmt}": f"old {fmt}".encode() for fmt in ("csv", "json", "xlsx")}
    for name, content in before.items():
        (tmp_path / name).write_bytes(content)
    real_replace = module.os.replace

    def replace(source, destination):
        result = real_replace(source, destination)
        if Path(destination).name == f"{STEM}.csv" and Path(source).suffix == ".tmp":
            cancel.set()
        return result

    monkeypatch.setattr(module.os, "replace", replace)
    with pytest.raises(CancelledError):
        module.export_contacts(_records(), "all", tmp_path, overwrite=True, cancel=cancel)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


def test_write_permission_failure_does_not_move_any_destination(tmp_path, monkeypatch):
    module = _export_module()
    before = {f"{STEM}.{fmt}": b"original" for fmt in ("csv", "json", "xlsx")}
    for name, content in before.items():
        (tmp_path / name).write_bytes(content)

    def save(*args, **kwargs):
        raise PermissionError("injected workbook write failure")

    monkeypatch.setattr(ExcelWriter, "write_data", save)
    with pytest.raises(PermissionError, match="write failure"):
        module.export_contacts(_records(), "all", tmp_path, overwrite=True)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


def test_new_file_appearing_during_staging_is_not_overwritten(tmp_path):
    target = tmp_path / "late.json"

    def report(done, total):
        if done == total:
            target.write_bytes(b"late original")

    with pytest.raises(FileExistsError):
        _export_module().export_contacts(_records(), "json", target, progress=report)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {target.name: b"late original"}


def test_invalid_format_does_not_create_a_target(tmp_path):
    with pytest.raises(ValueError):
        _export_module().export_contacts(_records(), "pdf", tmp_path / "new")
    assert list(tmp_path.iterdir()) == []


def test_xlsx_rejects_text_that_excel_would_silently_truncate(tmp_path):
    target = tmp_path / "out.xlsx"
    target.write_bytes(b"original")
    records = [dict.fromkeys(FIELDS, "x" * 32768)]
    with pytest.raises(ValueError, match="32767|too long|length"):
        _export_module().export_contacts(records, "xlsx", target, overwrite=True)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {"out.xlsx": b"original"}


@pytest.mark.filterwarnings("error::pytest.PytestUnraisableExceptionWarning")
def test_xlsx_cancellation_checks_during_archive_serialization(tmp_path, monkeypatch):
    module = _export_module()
    cancel = Event()
    target = tmp_path / "out.xlsx"
    target.write_bytes(b"original")
    real_write = ZipFile.writestr
    completed_entries = []

    def write(archive, name, content, *args, **kwargs):
        result = real_write(archive, name, content, *args, **kwargs)
        completed_entries.append(str(name))
        cancel.set()
        return result

    monkeypatch.setattr(ZipFile, "writestr", write)
    with pytest.raises(CancelledError):
        module.export_contacts(_records(), "xlsx", target, overwrite=True, cancel=cancel)
    assert len(completed_entries) == 1
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {"out.xlsx": b"original"}


def test_no_overwrite_is_exclusive_even_after_the_last_existence_check(tmp_path, monkeypatch):
    module = _export_module()
    target = tmp_path / "race.json"
    real_check = module._check_destination
    checks = 0

    def check(destination, overwrite):
        nonlocal checks
        real_check(destination, overwrite)
        checks += 1
        if checks == 3:
            target.write_bytes(b"concurrent original")

    monkeypatch.setattr(module, "_check_destination", check)
    with pytest.raises(FileExistsError):
        module.export_contacts(_records(), "json", target)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {target.name: b"concurrent original"}


def test_progress_callback_error_cleans_stages_without_touching_originals(tmp_path):
    target = tmp_path / "out.csv"
    target.write_bytes(b"original")

    def report(done, total):
        if done:
            raise RuntimeError("progress handler failed")

    with pytest.raises(RuntimeError, match="progress handler"):
        _export_module().export_contacts(_records(), "csv", target, overwrite=True, progress=report)
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == {target.name: b"original"}


def test_directory_is_never_replaced_even_when_overwrite_is_allowed(tmp_path):
    target = tmp_path / "out.json"
    target.mkdir()
    unrelated = target / "keep.txt"
    unrelated.write_bytes(b"unrelated")
    with pytest.raises(IsADirectoryError):
        _export_module().export_contacts(_records(), "json", target, overwrite=True)
    assert unrelated.read_bytes() == b"unrelated"
    assert list(tmp_path.iterdir()) == [target]


def test_json_preserves_numeric_zero_as_text_and_empty_cells(tmp_path):
    target = tmp_path / "out.json"
    _export_module().export_contacts([{"phone": 0, "username": "0", "remark": None}], "json", target)
    assert json.loads(target.read_text(encoding="utf-8")) == [{"nick_name": "", "remark": "", "phone": "0", "username": "0", "alias": "", "description": ""}]


@pytest.mark.parametrize("outcome", ("success", "write_error", "cancel"))
def test_xlsx_never_spools_contact_xml_outside_selected_staging_area(tmp_path, monkeypatch, outcome):
    module = _export_module()
    selected = tmp_path / "selected"
    selected.mkdir()
    target = selected / "contacts.xlsx"
    target.write_bytes(b"original")
    cancel = Event()
    created_paths = []
    real_named_temporary = worksheet_writer.NamedTemporaryFile
    real_mkstemp = tempfile.mkstemp
    original_tempdir = tempfile.gettempdir()

    def named_temporary(*args, **kwargs):
        result = real_named_temporary(*args, **kwargs)
        created_paths.append(Path(result.name))
        return result

    def mkstemp(*args, **kwargs):
        result = real_mkstemp(*args, **kwargs)
        created_paths.append(Path(result[1]))
        return result

    def report(done, total):
        if outcome == "cancel" and done >= 128:
            cancel.set()

    monkeypatch.setattr(worksheet_writer, "NamedTemporaryFile", named_temporary)
    monkeypatch.setattr(tempfile, "mkstemp", mkstemp)
    if outcome == "write_error":
        def write_error(*args, **kwargs):
            raise PermissionError("injected serialization error")
        monkeypatch.setattr(ExcelWriter, "write_data", write_error)
    records = _records() * 300
    records[0] = {**records[0], "description": "private\rcontact"}
    if outcome == "success":
        module.export_contacts(records, "xlsx", target, overwrite=True, cancel=cancel, progress=report)
        workbook = load_workbook(target)
        assert workbook.active.cell(2, 6).value == "private\rcontact"
        workbook.close()
    else:
        expected = PermissionError if outcome == "write_error" else CancelledError
        with pytest.raises(expected):
            module.export_contacts(records, "xlsx", target, overwrite=True, cancel=cancel, progress=report)
        assert target.read_bytes() == b"original"
    assert created_paths
    assert all(path.parent == selected for path in created_paths), created_paths
    assert tempfile.gettempdir() == original_tempdir
    assert list(selected.iterdir()) == [target]
