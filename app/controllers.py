from __future__ import annotations

import csv
import copy
import json
import os
import platform
import sys
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from PySide6.QtCore import QObject, Property, QSettings, QTimer, Signal, Slot

from .agent.client import AgentClient
from .agent.diagnostics import default_log_dir, redact_identifier
from .agent.workflows import normalize_identity
from .models import extract_greeting_name
from .task_models import (
    FriendImportModel,
    RuntimeLogModel,
    TaskDisplayItem,
    TaskItemModel,
)


MESSAGE_STEP_LABELS = {
    "window_bound": "已绑定微信窗口",
    "search_ready": "搜索入口已就绪",
    "target_selected": "已选择目标",
    "target_verified": "目标校验通过",
    "composer_ready": "输入框已就绪",
    "content_inserted": "内容已写入",
    "send_triggered": "已触发发送",
    "send_verified": "发送结果已确认",
}
FRIEND_STEP_LABELS = {
    "window_bound": "已绑定微信窗口",
    "add_friend_window_ready": "添加好友窗口已就绪",
    "account_inserted": "账号已写入",
    "account_searched": "已搜索账号",
    "profile_verified": "资料核对通过",
    "request_form_ready": "申请窗口已就绪",
    "fields_verified": "申请内容已核对",
    "preflight_completed": "表单预检已完成",
    "submit_triggered": "已点击确定",
    "submit_verified": "提交结果已确认",
}

FAILED_STEP_LABELS = {
    "window_bound": "绑定微信窗口失败",
    "search_ready": "搜索入口准备失败",
    "target_selected": "选择目标失败",
    "target_verified": "目标校验失败",
    "composer_ready": "消息输入框准备失败",
    "content_inserted": "写入消息内容失败",
    "send_triggered": "触发发送失败",
    "send_verified": "发送结果核对失败",
    "add_friend_window_ready": "打开添加好友窗口失败",
    "account_inserted": "写入账号失败",
    "account_searched": "搜索账号失败",
    "profile_verified": "资料核对失败",
    "request_form_ready": "打开申请窗口失败",
    "fields_verified": "申请内容核对失败",
    "preflight_completed": "表单预检失败",
    "submit_triggered": "点击确定失败",
    "submit_verified": "提交结果核对失败",
}

ERROR_RECOVERY_HINTS = {
    "TRANSIENT_UI": "微信界面暂时未就绪；可在任务结束后安全重试本条。",
    "STALE_ELEMENT": "微信控件已变化；Agent 已刷新会话，可安全重试本条。",
    "WECHAT_UNRESPONSIVE": "微信窗口无响应；请先检测恢复，仍无响应时再重启微信。",
    "UNSUPPORTED_VERSION": "当前微信版本未通过验证，不能执行自动化。",
    "GATE_SAFETY": "微信自动化安全门禁校验失败；任务已停止，请导出诊断包。",
    "UIA_TREE_NOT_READY_AFTER_REFRESH": (
        "可访问性广播已刷新，但微信仍未生成完整 UIA 树；"
        "请导出诊断包检查重复 Agent 或微信可访问性提供程序。"
    ),
    "RISK_CONTROL": "检测到验证码、频率或账号限制，任务已停止，请勿立即重试。",
    "TARGET_NOT_FOUND": "未找到精确目标；请检查微信名后再开始新任务。",
    "TARGET_NOT_UNIQUE": "搜索结果不唯一；请改用可唯一识别的微信号。",
    "RESULT_UNKNOWN": "动作已经触发但结果无法确认；为防止重复，本条不能重试。",
    "RESULT_VERIFICATION_FAILED": "动作已经触发但结果核对失败；本条不能重试。",
    "SUBMIT_NOT_TRIGGERED": "未能安全命中“确定”按钮，本条未提交；请检查微信窗口后重新开始。",
    "DESTRUCTIVE_BOUNDARY_UNKNOWN": "动作已越过发送边界但结果未知；不会自动重发，请人工核对微信记录。",
    "AUTOMATION_ERROR": "自动化步骤失败，请导出诊断包后检查具体原因。",
}

HEALTH_FAILURE_REASONS = frozenset(ERROR_RECOVERY_HINTS) | {
    "CLEANUP_FAILED", "HEALTH_CHECK_FAILED", "WINDOW_BLOCKED", "WINDOW_DISABLED",
}


def _redact_diagnostic(value: Any) -> Any:
    if isinstance(value, dict):
        private_fields = {"detail", "target", "account", "message", "greeting", "remark", "title", "name", "filePaths"}
        return {
            key: redact_identifier(str(item)) if key in private_fields and item else _redact_diagnostic(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [_redact_diagnostic(item) for item in value]
    return value


class MessageController(QObject):
    recipientsTextChanged = Signal(str)
    templateTextChanged = Signal(str)
    useForwardChanged = Signal(bool)
    filePathsChanged = Signal(list)
    intervalMinChanged = Signal(float)
    intervalMaxChanged = Signal(float)
    previewChanged = Signal()

    def __init__(self, settings: QSettings, parent=None):
        super().__init__(parent)
        self._settings = settings
        self._recipients = ""
        self._template = ""
        self._use_forward = False
        self._files: list[str] = []
        self._interval_min = float(settings.value("message/intervalMin", 2.0))
        self._interval_max = float(settings.value("message/intervalMax", 3.0))

    def _get_recipients(self):
        return self._recipients

    def _set_recipients(self, value):
        value = str(value)
        if value == self._recipients:
            return
        self._recipients = value
        self.recipientsTextChanged.emit(value)
        self.previewChanged.emit()

    recipientsText = Property(
        str, _get_recipients, _set_recipients, notify=recipientsTextChanged
    )

    def _get_template(self):
        return self._template

    def _set_template(self, value):
        value = str(value)
        if value == self._template:
            return
        self._template = value
        self.templateTextChanged.emit(value)
        self.previewChanged.emit()

    templateText = Property(str, _get_template, _set_template, notify=templateTextChanged)

    def _get_forward(self):
        return self._use_forward

    def _set_forward(self, value):
        value = bool(value)
        if value == self._use_forward:
            return
        self._use_forward = value
        self.useForwardChanged.emit(value)

    useForward = Property(bool, _get_forward, _set_forward, notify=useForwardChanged)

    @Property(list, notify=filePathsChanged)
    def filePaths(self):
        return list(self._files)

    def _get_interval_min(self):
        return self._interval_min

    def _set_interval_min(self, value):
        value = max(0.0, min(float(value), self._interval_max))
        if value == self._interval_min:
            return
        self._interval_min = value
        self._settings.setValue("message/intervalMin", value)
        self.intervalMinChanged.emit(value)

    intervalMin = Property(
        float, _get_interval_min, _set_interval_min, notify=intervalMinChanged
    )

    def _get_interval_max(self):
        return self._interval_max

    def _set_interval_max(self, value):
        value = max(self._interval_min, float(value))
        if value == self._interval_max:
            return
        self._interval_max = value
        self._settings.setValue("message/intervalMax", value)
        self.intervalMaxChanged.emit(value)

    intervalMax = Property(
        float, _get_interval_max, _set_interval_max, notify=intervalMaxChanged
    )

    @Property(int, notify=recipientsTextChanged)
    def recipientCount(self):
        return len(self.recipients())

    @Property(str, notify=previewChanged)
    def previewTarget(self):
        recipients = self.recipients()
        return recipients[0] if recipients else ""

    @Property(str, notify=previewChanged)
    def previewMessage(self):
        target = self.previewTarget
        if not target:
            return ""
        try:
            return self._template.format(name=extract_greeting_name(target))
        except (KeyError, ValueError):
            return self._template

    def recipients(self) -> list[str]:
        result = []
        seen = set()
        for line in self._recipients.splitlines():
            target = line.strip()
            identity = normalize_identity(target)
            if not identity or identity in seen:
                continue
            seen.add(identity)
            result.append(target)
        return result

    def build_items(self) -> list[dict[str, str]]:
        items = []
        for target in self.recipients():
            greeting = extract_greeting_name(target)
            try:
                message = self._template.format(name=greeting)
            except (KeyError, ValueError):
                message = self._template
            items.append(
                {
                    "itemId": uuid.uuid4().hex,
                    "target": target,
                    "message": message,
                }
            )
        return items

    @Slot(str)
    def addFile(self, file_url: str) -> None:
        path = file_url.replace("file:///", "")
        if sys.platform == "win32":
            path = path.lstrip("/")
        if path and path not in self._files:
            self._files.append(path)
            self.filePathsChanged.emit(list(self._files))
            self.previewChanged.emit()

    @Slot(int)
    def removeFile(self, index: int) -> None:
        if 0 <= index < len(self._files):
            self._files.pop(index)
            self.filePathsChanged.emit(list(self._files))
            self.previewChanged.emit()


class FriendController(QObject):
    defaultGreetingChanged = Signal(str)
    defaultRemarkChanged = Signal(str)
    intervalMinChanged = Signal(float)
    intervalMaxChanged = Signal(float)

    DEFAULT_GREETING = "你好，我是五阿哥，方便认识一下吗？"
    DEFAULT_REMARK = "新联系人"

    def __init__(self, settings: QSettings, parent=None):
        super().__init__(parent)
        self._settings = settings
        self._model = FriendImportModel(self)
        self._default_greeting = str(
            settings.value("friends/defaultGreeting", self.DEFAULT_GREETING)
        )
        self._default_remark = str(
            settings.value("friends/defaultRemark", self.DEFAULT_REMARK)
        )
        raw_interval_min = float(settings.value("friends/intervalMin", 15.0))
        raw_interval_max = float(settings.value("friends/intervalMax", 30.0))
        bounded_min = max(1.0, min(300.0, raw_interval_min))
        bounded_max = max(1.0, min(300.0, raw_interval_max))
        self._interval_min = min(bounded_min, bounded_max)
        self._interval_max = max(bounded_min, bounded_max)
        if self._interval_min != raw_interval_min:
            settings.setValue("friends/intervalMin", self._interval_min)
        if self._interval_max != raw_interval_max:
            settings.setValue("friends/intervalMax", self._interval_max)

    @Property(QObject, constant=True)
    def model(self):
        return self._model

    def _get_greeting(self):
        return self._default_greeting

    def _set_greeting(self, value):
        value = str(value)
        if value == self._default_greeting:
            return
        self._default_greeting = value
        self._settings.setValue("friends/defaultGreeting", value)
        self.defaultGreetingChanged.emit(value)

    defaultGreeting = Property(
        str, _get_greeting, _set_greeting, notify=defaultGreetingChanged
    )

    def _get_remark(self):
        return self._default_remark

    def _set_remark(self, value):
        value = str(value)
        if value == self._default_remark:
            return
        self._default_remark = value
        self._settings.setValue("friends/defaultRemark", value)
        self.defaultRemarkChanged.emit(value)

    defaultRemark = Property(str, _get_remark, _set_remark, notify=defaultRemarkChanged)

    def _get_interval_min(self):
        return self._interval_min

    def _set_interval_min(self, value):
        value = max(1.0, min(300.0, float(value)))
        value = min(value, self._interval_max)
        if value == self._interval_min:
            return
        self._interval_min = value
        self._settings.setValue("friends/intervalMin", value)
        self.intervalMinChanged.emit(value)

    intervalMin = Property(
        float, _get_interval_min, _set_interval_min, notify=intervalMinChanged
    )

    def _get_interval_max(self):
        return self._interval_max

    def _set_interval_max(self, value):
        value = max(1.0, min(300.0, float(value)))
        value = max(value, self._interval_min)
        if value == self._interval_max:
            return
        self._interval_max = value
        self._settings.setValue("friends/intervalMax", value)
        self.intervalMaxChanged.emit(value)

    intervalMax = Property(
        float, _get_interval_max, _set_interval_max, notify=intervalMaxChanged
    )

    @Slot(str, result=bool)
    def importFile(self, path: str) -> bool:
        return self._model.importFile(path)

    def build_items(self) -> list[dict[str, Any]]:
        return self._model.selected_payload(
            self._default_greeting, self._default_remark
        )

    @Slot(str, result=bool)
    def createTemplate(self, file_url: str) -> bool:
        path = file_url.replace("file:///", "")
        if sys.platform == "win32":
            path = path.lstrip("/")
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "好友导入"
            worksheet.append(["账号", "打招呼语", "备注"])
            worksheet.append(["18896904196", "你好，方便认识一下吗？", "示例联系人"])
            worksheet.freeze_panes = "A2"
            worksheet.column_dimensions["A"].width = 24
            worksheet.column_dimensions["B"].width = 42
            worksheet.column_dimensions["C"].width = 24
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="07C160")
                cell.alignment = Alignment(vertical="center")
            worksheet.row_dimensions[1].height = 24
            os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
            workbook.save(path)
            workbook.close()
            return True
        except Exception:
            return False


class SettingsController(QObject):
    unknownPolicyChanged = Signal(str)
    agentRestartLimitChanged = Signal(int)
    wechatRecoveryModeChanged = Signal(str)
    loginTimeoutChanged = Signal(int)
    appearanceChanged = Signal()

    def __init__(self, owner, settings: QSettings, parent=None):
        super().__init__(parent)
        self._owner = owner
        self._settings = settings
        self._unknown_policy = str(
            settings.value("task/unknownPolicy", "continue")
        )
        self._agent_restart_limit = int(settings.value("recovery/agentRestarts", 2))
        self._wechat_mode = str(settings.value("recovery/wechatMode", "confirm"))
        self._login_timeout = int(settings.value("recovery/loginTimeout", 90))
        owner.isDarkChanged.connect(lambda _value: self.appearanceChanged.emit())
        owner.glassEnabledChanged.connect(lambda _value: self.appearanceChanged.emit())
        owner.glassOpacityChanged.connect(lambda _value: self.appearanceChanged.emit())

    def _get_unknown_policy(self):
        return self._unknown_policy

    def _set_unknown_policy(self, value):
        value = str(value)
        if value not in {"continue", "stop"} or value == self._unknown_policy:
            return
        self._unknown_policy = value
        self._settings.setValue("task/unknownPolicy", value)
        self.unknownPolicyChanged.emit(value)

    unknownPolicy = Property(
        str, _get_unknown_policy, _set_unknown_policy, notify=unknownPolicyChanged
    )

    def _get_restart_limit(self):
        return self._agent_restart_limit

    def _set_restart_limit(self, value):
        value = max(0, min(2, int(value)))
        if value == self._agent_restart_limit:
            return
        self._agent_restart_limit = value
        self._settings.setValue("recovery/agentRestarts", value)
        self.agentRestartLimitChanged.emit(value)

    agentRestartLimit = Property(
        int,
        _get_restart_limit,
        _set_restart_limit,
        notify=agentRestartLimitChanged,
    )

    def _get_wechat_mode(self):
        return self._wechat_mode

    def _set_wechat_mode(self, value):
        value = str(value)
        if value not in {"confirm", "manual", "silent"} or value == self._wechat_mode:
            return
        self._wechat_mode = value
        self._settings.setValue("recovery/wechatMode", value)
        self.wechatRecoveryModeChanged.emit(value)

    wechatRecoveryMode = Property(
        str,
        _get_wechat_mode,
        _set_wechat_mode,
        notify=wechatRecoveryModeChanged,
    )

    def _get_login_timeout(self):
        return self._login_timeout

    def _set_login_timeout(self, value):
        value = max(30, min(300, int(value)))
        if value == self._login_timeout:
            return
        self._login_timeout = value
        self._settings.setValue("recovery/loginTimeout", value)
        self.loginTimeoutChanged.emit(value)

    loginTimeout = Property(
        int, _get_login_timeout, _set_login_timeout, notify=loginTimeoutChanged
    )

    isDark = Property(
        bool,
        lambda self: self._owner.isDark,
        lambda self, value: setattr(self._owner, "isDark", bool(value)),
        notify=appearanceChanged,
    )
    glassEnabled = Property(
        bool,
        lambda self: self._owner.glassEnabled,
        lambda self, value: setattr(self._owner, "glassEnabled", bool(value)),
        notify=appearanceChanged,
    )
    glassOpacity = Property(
        int,
        lambda self: self._owner.glassOpacity,
        lambda self, value: setattr(self._owner, "glassOpacity", int(value)),
        notify=appearanceChanged,
    )


class AgentController(QObject):
    stateChanged = Signal()
    inspectionChanged = Signal()
    notificationReceived = Signal(str, object)
    errorOccurred = Signal(str)
    helloReceived = Signal(object)
    connectionLost = Signal()
    replyReceived = Signal(int, object)
    rpcErrorReceived = Signal(int, int, str)

    def __init__(self, client: AgentClient | None = None, parent=None):
        super().__init__(parent)
        self._client = client or AgentClient(parent=self)
        self._state = str(getattr(self._client, "state", "stopped"))
        self._connected = bool(getattr(self._client, "connected", False))
        self._process_detected = False
        self._version_supported = False
        self._session_ready = False
        self._window_responsive = False
        self._window_enabled = False
        self._blocking_window: Any = None
        self._health_snapshot: dict[str, Any] = {}
        self._health_sequence = -1
        self._health_instance_id = ""
        self._retired_instance_ids: set[str] = set()
        self._friend_submit_enabled: bool | None = None
        self._build_fingerprint = ""
        self._session_generation = 0
        self._degraded_reason = ""
        self._wechat_connected = False
        self._wechat_supported = False
        self._uia_ready = False
        self._restorable = False
        self._wechat_version = ""
        self._detail = "等待检测微信"
        self._inspect_request_id = 0
        self._client.stateChanged.connect(self._on_state)
        self._client.connectedChanged.connect(self._on_connected)
        self._client.replyReceived.connect(self._on_reply)
        self._client.rpcError.connect(self._on_rpc_error)
        self._client.notificationReceived.connect(self._on_notification)
        self._client.processError.connect(self.errorOccurred)
        if hasattr(self._client, "helloReceived"):
            self._client.helloReceived.connect(self._on_hello)

    @Property("QVariant", notify=inspectionChanged)
    def friendSubmitEnabled(self):
        return self._friend_submit_enabled

    @Property(str, notify=inspectionChanged)
    def buildFingerprint(self):
        return self._build_fingerprint

    @Slot(object)
    def _on_hello(self, result: Any) -> None:
        capabilities = result.get("capabilities", {}) if isinstance(result, dict) else {}
        enabled = capabilities.get("friendSubmitEnabled") if isinstance(capabilities, dict) else None
        self._friend_submit_enabled = enabled if type(enabled) is bool else None
        build = result.get("build", {}) if isinstance(result, dict) else {}
        fingerprint = build.get("buildFingerprint", "") if isinstance(build, dict) else ""
        self._build_fingerprint = fingerprint if isinstance(fingerprint, str) else ""
        self.inspectionChanged.emit()
        self.helloReceived.emit(result)

    @Property(str, notify=stateChanged)
    def state(self):
        return self._state

    @Property(bool, notify=stateChanged)
    def connected(self):
        return self._connected

    @Property(bool, notify=inspectionChanged)
    def wechatConnected(self):
        return self._wechat_connected

    @Property(bool, notify=inspectionChanged)
    def processDetected(self):
        return self._process_detected

    @Property(bool, notify=inspectionChanged)
    def wechatSupported(self):
        return self._wechat_supported

    @Property(bool, notify=inspectionChanged)
    def versionSupported(self):
        return self._version_supported

    @Property(bool, notify=inspectionChanged)
    def uiaReady(self):
        return self._uia_ready

    @Property(bool, notify=inspectionChanged)
    def sessionReady(self):
        return self._session_ready

    @Property(bool, notify=inspectionChanged)
    def windowResponsive(self):
        return self._window_responsive

    @Property(bool, notify=inspectionChanged)
    def windowEnabled(self):
        return self._window_enabled

    @Property("QVariant", notify=inspectionChanged)
    def blockingWindow(self):
        return copy.deepcopy(self._blocking_window)

    @Property("QVariantMap", notify=inspectionChanged)
    def healthSnapshot(self):
        return copy.deepcopy(self._health_snapshot)

    @Property(str, notify=inspectionChanged)
    def reasonCode(self):
        if not self._connected:
            return "AGENT_DISCONNECTED"
        return str(self._health_snapshot.get("reasonCode", self._degraded_reason))

    @Property(int, notify=inspectionChanged)
    def sessionGeneration(self):
        return self._session_generation

    @Property(str, notify=inspectionChanged)
    def degradedReason(self):
        return self._degraded_reason

    @Property(bool, notify=inspectionChanged)
    def restorable(self):
        return self._restorable

    @Property(str, notify=inspectionChanged)
    def wechatVersion(self):
        return self._wechat_version

    @Property(str, notify=inspectionChanged)
    def detail(self):
        return self._detail

    @Property(bool, notify=inspectionChanged)
    def automationReady(self):
        return self.canStartTask and self._session_ready

    @Property(bool, notify=inspectionChanged)
    def canStartTask(self):
        return (
            self._connected
            and self._process_detected
            and self._version_supported
            and self._window_responsive
            and self._window_enabled
            and not self._blocking_window
            and self.reasonCode not in HEALTH_FAILURE_REASONS
            and self._degraded_reason not in HEALTH_FAILURE_REASONS
            and (self._session_ready or self._restorable)
        )

    @Slot()
    def start(self) -> None:
        self._client.start()

    @Slot()
    def restart(self) -> None:
        restart = getattr(self._client, "restart", None)
        if restart is not None:
            restart()
            return
        self._client.close()
        self._client.start()

    def recoverySnapshot(self) -> dict[str, Any] | None:
        snapshot = getattr(self._client, "recovery_snapshot", None)
        if not callable(snapshot):
            return None
        try:
            value = snapshot()
        except Exception:
            return None
        return dict(value) if isinstance(value, dict) else None

    @Slot(result=int)
    def inspect(self) -> int:
        if not self._connected:
            return 0
        self._inspect_request_id = self._client.call("wechat.inspect")
        return self._inspect_request_id

    @Slot(object)
    def applyInspection(self, result: dict[str, Any]) -> None:
        if not self._connected or not isinstance(result, dict):
            return
        instance = str(result.get("agentInstanceId", ""))
        sequence = result.get("sequence")
        if instance in self._retired_instance_ids:
            return
        if instance and isinstance(sequence, int) and not isinstance(sequence, bool):
            if instance == self._health_instance_id and sequence <= self._health_sequence:
                return
            if self._health_instance_id and instance != self._health_instance_id:
                self._retired_instance_ids.add(self._health_instance_id)
            self._health_instance_id = instance
            self._health_sequence = sequence
        elif self._health_instance_id:
            # An unversioned reply cannot supersede an ordered health snapshot.
            return
        self._health_snapshot = copy.deepcopy(result)
        self._process_detected = bool(
            result.get("processDetected", result.get("connected", False))
        )
        self._version_supported = bool(
            result.get("versionSupported", result.get("supported", False))
        )
        self._session_ready = bool(
            result.get("sessionReady", result.get("uiaReady", False))
        )
        self._window_responsive = bool(
            result.get("windowResponsive", self._process_detected)
        )
        self._window_enabled = bool(result.get("windowEnabled", self._process_detected))
        self._blocking_window = copy.deepcopy(result.get("blockingWindow"))
        self._session_generation = int(result.get("sessionGeneration", 0) or 0)
        self._degraded_reason = str(result.get("degradedReason", result.get("reasonCode", "")))
        self._wechat_connected = self._process_detected
        self._wechat_supported = self._version_supported
        self._uia_ready = self._session_ready
        self._restorable = bool(result.get("restorable", False))
        self._wechat_version = str(result.get("version", ""))
        self._detail = str(result.get("detail", ""))
        self.inspectionChanged.emit()

    def call(self, method: str, params: dict[str, Any] | None = None) -> int:
        if method == "task.start":
            self._client.grant_foreground_permission()
        return self._client.call(method, params or {})

    @Slot(str, object)
    def _on_notification(self, method: str, params: Any) -> None:
        if method == "agent.status" and isinstance(params, dict):
            snapshot = params.get("health")
            if isinstance(snapshot, dict):
                self.applyInspection(snapshot)
        self.notificationReceived.emit(method, params)

    @Slot(str)
    def _on_state(self, state: str) -> None:
        self._state = state
        self.stateChanged.emit()

    @Slot(bool)
    def _on_connected(self, connected: bool) -> None:
        was_connected = self._connected
        self._connected = connected
        if not connected:
            self._inspect_request_id = 0
            self._friend_submit_enabled = None
            self._wechat_connected = False
            self._wechat_supported = False
            self._uia_ready = False
            self._process_detected = False
            self._version_supported = False
            self._session_ready = False
            self._window_responsive = False
            self._window_enabled = False
            self._blocking_window = None
            # Keep the last observed generation and snapshot for diagnostics only.
            self._degraded_reason = "AGENT_DISCONNECTED"
            self._restorable = False
            self._detail = "Agent 已断开"
        self.stateChanged.emit()
        self.inspectionChanged.emit()
        if connected:
            self.inspect()
        elif was_connected:
            self.connectionLost.emit()

    @Slot(int, object)
    def _on_reply(self, request_id: int, result: Any) -> None:
        if self._inspect_request_id and request_id == self._inspect_request_id and isinstance(result, dict):
            self._inspect_request_id = 0
            self.applyInspection(result)
        self.replyReceived.emit(request_id, result)

    @Slot(int, int, str)
    def _on_rpc_error(self, request_id: int, code: int, message: str) -> None:
        self.rpcErrorReceived.emit(request_id, code, message)
        self.errorOccurred.emit(message)

    def close(self) -> None:
        shutdown = getattr(self._client, "shutdown", None)
        if shutdown is not None:
            shutdown()
        else:
            if self._connected:
                try:
                    self._client.call("agent.shutdown")
                except Exception:
                    pass
            self._client.close()


class TaskController(QObject):
    phaseChanged = Signal()
    activeChanged = Signal()
    kindChanged = Signal()
    progressChanged = Signal()
    currentStepChanged = Signal()
    errorChanged = Signal()
    recoveryRequiredChanged = Signal()
    executionStateChanged = Signal()
    acceptanceStateChanged = Signal()

    def __init__(
        self,
        agent: AgentController,
        message: MessageController,
        friends: FriendController,
        settings: SettingsController,
        parent=None,
    ):
        super().__init__(parent)
        self._agent = agent
        self._message = message
        self._friends = friends
        self._settings = settings
        self._items = TaskItemModel(self)
        self._logs = RuntimeLogModel(self)
        self._phase = "idle"
        self._active = False
        self._task_id = ""
        self._kind = ""
        self._done = 0
        self._total = 0
        self._current_step = ""
        self._current_outcome = "pending"
        self._current_detail = ""
        self._current_error_code = ""
        self._retry_attempt = 1
        self._retry_max_attempts = 1
        self._retry_level = "none"
        self._recoverable = False
        self._destructive_boundary_crossed = False
        self._wechat_responsive = True
        self._retry_candidate: dict[str, Any] | None = None
        self._non_retryable_item_ids: set[str] = set()
        self._cleanup_result: dict[str, Any] = {}
        self._task_events: list[dict[str, Any]] = []
        self._finished_result: dict[str, Any] = {}
        self._acceptance_enabled = os.environ.get("WECHAT_COURIER_ACCEPTANCE") == "1"
        self._error = ""
        self._original_payload: dict[str, Any] | None = None
        self._completed_item_ids: set[str] = set()
        self._recovery_offset = 0
        self._restart_attempts = 0
        self._pending_resume = False
        self._recovery_required = False
        self._recovery_detail = ""
        self._wechat_restart_attempted = False
        self._pending_start_request_id = 0
        self._restart_scheduled = False
        self._agent_restart_timer = QTimer(self)
        self._agent_restart_timer.setSingleShot(True)
        self._agent_restart_timer.timeout.connect(self._perform_agent_restart)
        self._login_deadline = 0.0
        self._inspection_grace = QTimer(self)
        self._inspection_grace.setSingleShot(True)
        self._inspection_grace.setInterval(6_000)
        self._inspection_grace.timeout.connect(self._check_reconnect_inspection)
        self._recovery_poll = QTimer(self)
        self._recovery_poll.setInterval(2_000)
        self._recovery_poll.timeout.connect(self._poll_wechat_recovery)
        self._agent.notificationReceived.connect(self._on_notification)
        self._agent.errorOccurred.connect(self._set_error)
        self._agent.connectionLost.connect(self._on_connection_lost)
        self._agent.helloReceived.connect(self._on_agent_hello)
        self._agent.inspectionChanged.connect(self._resume_if_ready)
        self._agent.inspectionChanged.connect(self._sync_wechat_health)
        self._agent.replyReceived.connect(self._on_agent_reply)
        self._agent.rpcErrorReceived.connect(self._on_agent_rpc_error)
        self._agent.inspectionChanged.connect(self.executionStateChanged.emit)
        for signal in (
            self.activeChanged, self.phaseChanged, self.kindChanged, self.progressChanged, self.executionStateChanged,
            message.recipientsTextChanged, message.templateTextChanged, message.filePathsChanged,
            message.useForwardChanged, message.intervalMinChanged, message.intervalMaxChanged,
            friends.model.countsChanged, friends.defaultGreetingChanged, friends.defaultRemarkChanged,
            friends.intervalMinChanged, friends.intervalMaxChanged, settings.unknownPolicyChanged,
        ):
            signal.connect(lambda *_: self.acceptanceStateChanged.emit())

    @Property(bool, constant=True)
    def acceptanceEnabled(self):
        return self._acceptance_enabled

    def _acceptance_editor_state(self, kind: str) -> dict[str, Any]:
        if not self._acceptance_enabled:
            return {}
        editor = self._message if kind == "message_send" else self._friends
        options = {
            "intervalMin": editor.intervalMin,
            "intervalMax": editor.intervalMax,
            "unknownPolicy": self._settings.unknownPolicy,
            "filePaths": self._message.filePaths if kind == "message_send" else [],
            "useForward": self._message.useForward if kind == "message_send" else False,
        }
        return {
            "active": self._active,
            "friendSubmitEnabled": self._agent.friendSubmitEnabled,
            "kind": kind,
            "items": [{key: value for key, value in item.items() if key != "itemId"}
                      for item in editor.build_items()],
            "options": options,
        }

    @Property("QVariantMap", notify=acceptanceStateChanged)
    def acceptanceMessageState(self):
        return self._acceptance_editor_state("message_send")

    @Property("QVariantMap", notify=acceptanceStateChanged)
    def acceptanceFriendState(self):
        return self._acceptance_editor_state("friend_add")

    # Serialize before QVariantMap conversion, which maps None to QML undefined.
    @Property(str, notify=acceptanceStateChanged)
    def acceptanceMessageStateJson(self):
        return json.dumps(self.acceptanceMessageState, ensure_ascii=False)

    @Property(str, notify=acceptanceStateChanged)
    def acceptanceFriendStateJson(self):
        return json.dumps(self.acceptanceFriendState, ensure_ascii=False)

    @Property(str, notify=acceptanceStateChanged)
    def acceptanceTaskStateJson(self):
        return json.dumps(self.acceptanceTaskState, ensure_ascii=False)

    @Property("QVariantMap", notify=acceptanceStateChanged)
    def acceptanceTaskState(self):
        if not self._acceptance_enabled:
            return {}
        return {
            "taskId": self._task_id,
            "kind": self._kind,
            "active": self._active,
            "phase": self._phase,
            "outcome": self._finished_result.get("outcome", "error" if self._phase == "error" else "working"),
            "done": self._done,
            "total": self._total,
            "success": self.successCount,
            "error": sum(item.result == "error" for item in self._items._items),
            "unknown": self.unknownCount,
            "stopped": sum(item.result == "stopped" for item in self._items._items),
            "echoedItems": copy.deepcopy((self._original_payload or {}).get("items", [])),
            "cleanup": copy.deepcopy(self._finished_result.get("cleanup", {})),
            "health": copy.deepcopy(self._finished_result.get("health", {})),
            "buildFingerprint": self._agent.buildFingerprint,
            "events": [{key: event.get(key, "") for key in ("taskId", "itemId", "step", "outcome")}
                       for event in self._task_events if "itemId" in event and "step" in event],
        }

    @Property(QObject, constant=True)
    def items(self):
        return self._items

    @Property(QObject, constant=True)
    def runtimeLogs(self):
        return self._logs

    @Property(str, notify=phaseChanged)
    def phase(self):
        return self._phase

    @Property(bool, notify=activeChanged)
    def active(self):
        return self._active

    @Property(str, notify=kindChanged)
    def kind(self):
        return self._kind

    @Property(int, notify=progressChanged)
    def done(self):
        return self._done

    @Property(int, notify=progressChanged)
    def total(self):
        return self._total

    @Property(float, notify=progressChanged)
    def progress(self):
        return self._done / self._total if self._total else 0.0

    @Property(str, notify=currentStepChanged)
    def currentStepCode(self):
        return self._current_step

    @Property(str, notify=currentStepChanged)
    def currentStepLabel(self):
        labels = MESSAGE_STEP_LABELS if self._kind == "message_send" else FRIEND_STEP_LABELS
        if self._current_outcome in {"error", "unknown"}:
            return FAILED_STEP_LABELS.get(self._current_step, "自动化步骤失败")
        return labels.get(self._current_step, "等待任务开始")

    @Property(str, notify=executionStateChanged)
    def currentOutcome(self):
        return self._current_outcome

    @Property(str, notify=executionStateChanged)
    def currentDetail(self):
        return self._current_detail

    @Property(str, notify=executionStateChanged)
    def currentErrorCode(self):
        return self._current_error_code

    @Property(str, notify=executionStateChanged)
    def recoveryHint(self):
        return ERROR_RECOVERY_HINTS.get(self._current_error_code, "")

    @Property(int, notify=executionStateChanged)
    def retryAttempt(self):
        return self._retry_attempt

    @Property(int, notify=executionStateChanged)
    def retryMaxAttempts(self):
        return self._retry_max_attempts

    @Property(str, notify=executionStateChanged)
    def retryLevel(self):
        return self._retry_level

    @Property(bool, notify=executionStateChanged)
    def recoverable(self):
        return self._recoverable

    @Property(bool, notify=executionStateChanged)
    def destructiveBoundaryCrossed(self):
        return self._destructive_boundary_crossed

    @Property(bool, notify=executionStateChanged)
    def wechatResponsive(self):
        return self._wechat_responsive

    @Property(bool, notify=executionStateChanged)
    def safeRetryAvailable(self):
        return bool(
            not self._active
            and self._retry_candidate is not None
            and self._retry_candidate.get("itemId") not in self._non_retryable_item_ids
            and self._agent.canStartTask
        )

    @Property("QVariantMap", notify=executionStateChanged)
    def cleanupResult(self):
        return copy.deepcopy(self._cleanup_result)

    @Property(bool, notify=executionStateChanged)
    def cleanupFailed(self):
        return self._cleanup_result.get("success") is False

    @Property(bool, notify=executionStateChanged)
    def wechatRestartAvailable(self):
        return bool(
            not self._active
            and self._current_error_code == "WECHAT_UNRESPONSIVE"
            and not self._wechat_responsive
            and self._agent.connected
            and not self._wechat_restart_attempted
        )

    @Property(int, notify=progressChanged)
    def successCount(self):
        return sum(item.result == "success" for item in self._items._items)

    @Property(int, notify=progressChanged)
    def failureCount(self):
        return sum(item.result == "error" for item in self._items._items)

    @Property(int, notify=progressChanged)
    def unknownCount(self):
        return sum(item.result == "unknown" for item in self._items._items)

    @Property(str, notify=errorChanged)
    def error(self):
        return self._error

    @Property(bool, notify=recoveryRequiredChanged)
    def recoveryRequired(self):
        return self._recovery_required

    @Property(str, notify=recoveryRequiredChanged)
    def recoveryDetail(self):
        return self._recovery_detail

    def _set_recovery_required(self, required: bool, detail: str = "") -> None:
        changed = (
            required != self._recovery_required or detail != self._recovery_detail
        )
        self._recovery_required = required
        self._recovery_detail = detail
        if changed:
            self.recoveryRequiredChanged.emit()

    def _set_phase(self, phase: str) -> None:
        if phase != self._phase:
            self._phase = phase
            self.phaseChanged.emit()

    def _set_active(self, active: bool) -> None:
        if active != self._active:
            self._active = active
            self.activeChanged.emit()

    def _set_kind(self, kind: str) -> None:
        if kind != self._kind:
            self._kind = kind
            self.kindChanged.emit()

    @Slot(str)
    def _set_error(self, error: str) -> None:
        self._error = error
        self.errorChanged.emit()

    def _start(self, kind: str, items: list[dict[str, Any]], options: dict[str, Any]) -> bool:
        if self._active:
            self._set_error("已有任务正在执行")
            return False
        if not self._agent.canStartTask:
            self._set_error("仅支持已验证并已连接的微信 4.1.13.65")
            return False
        if not items:
            self._set_error("没有可执行的数据")
            return False
        self._task_id = uuid.uuid4().hex
        self._set_kind(kind)
        self._done = 0
        self._total = len(items)
        self._current_step = ""
        self._current_outcome = "pending"
        self._current_detail = ""
        self._current_error_code = ""
        self._retry_attempt = 1
        self._retry_max_attempts = 1
        self._retry_level = "none"
        self._recoverable = False
        self._destructive_boundary_crossed = False
        self._wechat_responsive = True
        self._retry_candidate = None
        self._non_retryable_item_ids.clear()
        self._cleanup_result = {}
        self._task_events.clear()
        self._finished_result = {}
        self._error = ""
        self._logs.clear()
        self._items.replace(
            TaskDisplayItem(
                target=item.get("target") or item.get("account", ""),
                item_id=item["itemId"],
            )
            for item in items
        )
        payload = {
            "taskId": self._task_id,
            "kind": kind,
            "items": items,
            "options": options,
        }
        self._original_payload = copy.deepcopy(payload)
        self._completed_item_ids.clear()
        self._recovery_offset = 0
        self._restart_attempts = 0
        self._pending_resume = False
        self._set_recovery_required(False)
        self._wechat_restart_attempted = False
        self._inspection_grace.stop()
        self._recovery_poll.stop()
        self._agent_restart_timer.stop()
        self._restart_scheduled = False
        try:
            self._pending_start_request_id = self._agent.call(
                "task.start", payload
            )
        except Exception as exc:
            self._pending_start_request_id = 0
            self._set_error(str(exc))
            return False
        self._set_active(True)
        self._set_phase("running")
        self.phaseChanged.emit()
        self.progressChanged.emit()
        self.currentStepChanged.emit()
        self.executionStateChanged.emit()
        return True

    @Slot()
    def _on_connection_lost(self) -> None:
        if not self._active:
            return
        recovery = self._agent.recoverySnapshot()
        if self._phase == "stopping":
            self._pending_start_request_id = 0
            self._pending_resume = False
            self._retry_candidate = None
            self._inspection_grace.stop()
            self._recovery_poll.stop()
            self._agent_restart_timer.stop()
            self._restart_scheduled = False
            self._set_recovery_required(False)
            if (
                recovery and recovery.get("taskId") == self._task_id
                and any(item.item_id == recovery.get("itemId") for item in self._items._items)
            ):
                self._mark_boundary_item_unknown(recovery)
            for item in self._items._items:
                if item.item_id in self._non_retryable_item_ids:
                    self._mark_boundary_item_unknown({
                        "itemId": item.item_id,
                        "boundary": "submit_triggered" if self._kind == "friend_add" else "send_triggered",
                    })
            self._set_error("停止任务时 Agent 已断开；未确认的结果与清理状态请人工核对，任务不会自动恢复")
            self._set_phase("error")
            self._set_active(False)
            self.executionStateChanged.emit()
            return
        if recovery and recovery.get("taskId") == self._task_id:
            self._mark_boundary_item_unknown(recovery)
            self._pending_resume = False
            self._set_active(False)
            self._set_phase("error")
            self._set_error(
                "Agent 在破坏性动作后中断，本条已标记为结果未知；当前批次不会自动恢复"
            )
            self._schedule_agent_restart(for_readiness=True)
            return
        if (
            self._original_payload
            and self._original_payload.get("options", {}).get("useForward")
        ):
            self._pending_resume = False
            self._set_active(False)
            self._set_phase("error")
            self._set_error(
                "合并转发任务已中断，为避免重复上传或转发，不会自动恢复"
            )
            self._schedule_agent_restart(for_readiness=True)
            return
        if self._restart_attempts >= self._settings.agentRestartLimit:
            self._set_active(False)
            self._set_phase("error")
            self._set_error("Agent 已断开，自动重启次数已用尽；任务已停止")
            return
        self._restart_attempts += 1
        self._pending_resume = False
        self._set_phase("recovering")
        self._set_error(
            f"Agent 连接中断，第 {self._restart_attempts} 次安全恢复即将开始"
        )
        self._schedule_agent_restart(for_readiness=False, increment=False)

    def _schedule_agent_restart(
        self,
        *,
        for_readiness: bool,
        increment: bool = True,
    ) -> None:
        if increment:
            if self._restart_attempts >= self._settings.agentRestartLimit:
                return
            self._restart_attempts += 1
        attempt = max(1, self._restart_attempts)
        delay = 2_000 if attempt == 1 else 5_000
        self._restart_scheduled = True
        self._agent_restart_timer.setProperty("forReadiness", bool(for_readiness))
        self._agent_restart_timer.start(delay)

    @Slot()
    def _perform_agent_restart(self) -> None:
        if not self._restart_scheduled:
            return
        self._restart_scheduled = False
        self._agent.restart()

    @Slot(object)
    def _on_agent_hello(self, hello: dict[str, Any]) -> None:
        recovery = hello.get("recovery") if isinstance(hello, dict) else None
        if not self._active:
            if recovery:
                self._agent.call("recovery.approve", {"decision": "discard"})
            return
        if self._phase != "recovering":
            return
        if recovery:
            if recovery.get("taskId") != self._task_id:
                self._set_active(False)
                self._set_phase("error")
                self._set_error("发现其他任务的未决安全记录，已停止自动恢复")
                return
            self._mark_boundary_item_unknown(recovery)
            self._agent.call(
                "recovery.approve", {"decision": "mark_unknown"}
            )
            self._pending_resume = False
            self._inspection_grace.stop()
            self._set_active(False)
            self._set_phase("error")
            self._set_error(
                "Agent 在破坏性动作后中断，本条已标记为结果未知；当前批次不会自动恢复"
            )
            return
        if not self._resume_agent_capabilities_match_payload():
            self._fail_recovery(
                "恢复好友任务失败：重启后的 Agent 未启用好友申请提交能力，任务已安全停止"
            )
            return
        self._pending_resume = True
        self._inspection_grace.start()
        self._resume_if_ready()

    def _resume_agent_capabilities_match_payload(self) -> bool:
        payload = self._original_payload
        if not isinstance(payload, dict) or payload.get("kind") != "friend_add":
            return True
        options = payload.get("options")
        submit_requested = (
            isinstance(options, dict)
            and options.get("submitFriendRequest") is True
        )
        return not submit_requested or self._agent.friendSubmitEnabled is True

    def _mark_boundary_item_unknown(self, recovery: dict[str, Any]) -> None:
        item_id = str(recovery.get("itemId", ""))
        if not item_id or item_id in self._completed_item_ids:
            return
        boundary = str(recovery.get("boundary", ""))
        step = "submit_verified" if boundary.startswith("submit") else "send_verified"
        detail = (
            "Agent 在提交后中断，本条标记为结果未知且不会再次提交"
            if self._kind == "friend_add"
            else "Agent 在发送后中断，本条标记为结果未知且不会自动重发"
        )
        event = {
            "taskId": self._task_id,
            "itemId": item_id,
            "step": step,
            "outcome": "unknown",
            "detail": detail,
            "done": min(self._total, self._done + 1),
            "total": self._total,
            "timestamp": str(recovery.get("timestamp", "")),
        }
        self._items.apply_event(event)
        self._logs.append_event(event)
        self._task_events.append(copy.deepcopy(event))
        self._non_retryable_item_ids.add(item_id)
        if self._retry_candidate and self._retry_candidate.get("itemId") == item_id:
            self._retry_candidate = None
        if self._kind == "friend_add":
            self._friends.model.apply_event(event)
        self._completed_item_ids.add(item_id)
        self._done = event["done"]
        self._current_step = step
        self._current_outcome = "unknown"
        self._current_detail = detail
        self._current_error_code = "DESTRUCTIVE_BOUNDARY_UNKNOWN"
        self._recoverable = False
        self._destructive_boundary_crossed = True
        self._wechat_responsive = True
        self.progressChanged.emit()
        self.currentStepChanged.emit()
        self.executionStateChanged.emit()

    @Slot()
    def _resume_if_ready(self) -> None:
        if not self._pending_resume or not self._active:
            return
        if not self._resume_agent_capabilities_match_payload():
            self._fail_recovery(
                "恢复好友任务失败：当前 Agent 未启用好友申请提交能力，任务已安全停止"
            )
            return
        if not self._agent.canStartTask or self._original_payload is None:
            return
        self._inspection_grace.stop()
        self._recovery_poll.stop()
        self._set_recovery_required(False)
        remaining = [
            item
            for item in self._original_payload["items"]
            if item["itemId"] not in self._completed_item_ids
        ]
        if not remaining:
            self._pending_resume = False
            self._set_active(False)
            self._set_phase("done")
            self._set_error("")
            return
        payload = copy.deepcopy(self._original_payload)
        payload["items"] = remaining
        self._recovery_offset = self._done
        try:
            self._pending_start_request_id = self._agent.call(
                "task.start", payload
            )
        except Exception as exc:
            self._pending_start_request_id = 0
            self._set_active(False)
            self._set_phase("error")
            self._set_error(f"恢复任务失败：{exc}")
            return
        self._pending_resume = False
        self._set_phase("running")
        self._set_error("")

    @Slot(int, object)
    def _on_agent_reply(self, request_id: int, result: Any) -> None:
        if request_id != self._pending_start_request_id:
            return
        self._pending_start_request_id = 0
        if isinstance(result, dict) and result.get("accepted", True):
            return
        self._set_active(False)
        self._set_phase("error")
        self._set_error("Agent 未接受任务")

    @Slot(int, int, str)
    def _on_agent_rpc_error(
        self, request_id: int, _code: int, message: str
    ) -> None:
        if request_id != self._pending_start_request_id:
            return
        self._pending_start_request_id = 0
        self._pending_resume = False
        self._set_active(False)
        self._set_phase("error")
        self._set_error(message)

    @Slot()
    def _check_reconnect_inspection(self) -> None:
        if not self._pending_resume or not self._active or self._agent.canStartTask:
            return
        if self._agent.degradedReason == "UIA_TREE_NOT_READY_AFTER_REFRESH":
            self._fail_recovery(
                "可访问性广播已刷新，但微信仍未生成完整 UIA 控件树；"
                "为避免无效重启循环，任务已停止，请导出诊断包"
            )
            return
        mode = self._settings.wechatRecoveryMode
        if mode == "manual":
            self._fail_recovery("微信 UIA 仍不可用，请手动重启微信后重新开始任务")
            return
        if mode == "silent":
            self._request_wechat_restart()
            return
        self._set_phase("awaiting_recovery")
        self._set_recovery_required(
            True,
            "Agent 已恢复，但微信 UIA 仍不可用。是否重启微信并等待重新登录？",
        )

    @Slot()
    def approveWechatRestart(self) -> None:
        if not self._recovery_required:
            return
        self._request_wechat_restart()

    def _request_wechat_restart(self) -> None:
        if self._wechat_restart_attempted or not self._active:
            return
        self._wechat_restart_attempted = True
        self._set_recovery_required(False)
        self._set_phase("waiting_login")
        self._set_error("正在重启微信并等待登录")
        self._login_deadline = time.monotonic() + self._settings.loginTimeout
        try:
            self._agent.call(
                "recovery.approve",
                {
                    "decision": "restart_wechat",
                    "loginTimeout": self._settings.loginTimeout,
                },
            )
        except Exception as exc:
            self._fail_recovery(f"请求重启微信失败：{exc}")
            return
        self._recovery_poll.start()

    @Slot()
    def stopRecovery(self) -> None:
        if not self._active:
            return
        self._agent.call("recovery.approve", {"decision": "stop"})
        self._fail_recovery("用户取消微信恢复，任务已安全停止")

    def _poll_wechat_recovery(self) -> None:
        if not self._active or not self._pending_resume:
            self._recovery_poll.stop()
            return
        if self._agent.canStartTask:
            self._resume_if_ready()
            return
        if time.monotonic() >= self._login_deadline:
            self._fail_recovery("等待微信登录超时，任务已停止")
            return
        self._agent.inspect()

    def _fail_recovery(self, detail: str) -> None:
        self._inspection_grace.stop()
        self._recovery_poll.stop()
        self._pending_resume = False
        self._set_recovery_required(False)
        self._set_active(False)
        self._set_phase("error")
        self._set_error(detail)

    @Slot(result=bool)
    def startMessage(self) -> bool:
        items = self._message.build_items()
        if not self._message.templateText and not self._message.filePaths:
            self._set_error("消息内容和附件不能同时为空")
            return False
        return self._start(
            "message_send",
            items,
            {
                "intervalMin": self._message.intervalMin,
                "intervalMax": self._message.intervalMax,
                "unknownPolicy": self._settings.unknownPolicy,
                "useForward": self._message.useForward,
                "filePaths": self._message.filePaths,
            },
        )

    @Slot(result=bool)
    def startFriends(self) -> bool:
        if self._active:
            self._set_error("已有任务正在执行")
            return False
        if (
            not self._acceptance_enabled
            and self._agent.friendSubmitEnabled is not True
        ):
            self._set_error(
                "当前 Agent 未启用好友申请提交能力，请重新启动或重新安装匹配版本"
            )
            return False
        return self._start(
            "friend_add",
            self._friends.build_items(),
            {
                "intervalMin": self._friends.intervalMin,
                "intervalMax": self._friends.intervalMax,
                "unknownPolicy": self._settings.unknownPolicy,
                "submitFriendRequest": not self._acceptance_enabled,
            },
        )

    @Slot()
    def pause(self) -> None:
        if self._active and self._phase == "running":
            self._agent.call("task.pause", {"taskId": self._task_id})
            self._set_phase("paused")

    @Slot()
    def resume(self) -> None:
        if self._active and self._phase == "paused":
            self._agent.call("task.resume", {"taskId": self._task_id})
            self._set_phase("running")

    @Slot()
    def stop(self) -> None:
        if self._active:
            self._agent.call("task.stop", {"taskId": self._task_id})
            self._set_phase("stopping")

    @Slot(str, object)
    def _on_notification(self, method: str, params: dict[str, Any]) -> None:
        if method == "task.event" and params.get("taskId") == self._task_id:
            params = dict(params)
            self._task_events.append(copy.deepcopy(params))
            if self._recovery_offset:
                params["done"] = min(
                    self._total,
                    self._recovery_offset + int(params.get("done", 0)),
                )
                params["total"] = self._total
            previous_done = self._done
            self._items.apply_event(params)
            self._logs.append_event(params)
            if self._kind == "friend_add":
                self._friends.model.apply_event(params)
            self._done = int(params.get("done", self._done))
            self._total = int(params.get("total", self._total))
            self._current_step = str(params.get("step", ""))
            self._current_outcome = str(params.get("outcome", "working"))
            self._current_detail = str(params.get("detail", ""))
            self._current_error_code = str(params.get("errorCode", ""))
            self._retry_attempt = int(params.get("attempt", 1) or 1)
            self._retry_max_attempts = int(params.get("maxAttempts", 1) or 1)
            self._retry_level = str(params.get("retryLevel", "none"))
            self._recoverable = bool(params.get("recoverable", False))
            self._destructive_boundary_crossed = bool(
                params.get("destructiveBoundaryCrossed", False)
            )
            self._wechat_responsive = bool(params.get("wechatResponsive", True))
            item_id = str(params.get("itemId", ""))
            if (
                self._destructive_boundary_crossed
                or self._current_outcome == "unknown"
                or self._current_step in {"send_triggered", "send_verified", "submit_verified"}
            ):
                self._non_retryable_item_ids.add(item_id)
            if self._retry_candidate and self._retry_candidate.get("itemId") == item_id:
                self._retry_candidate = None
            if (
                self._current_outcome == "error"
                and self._recoverable
                and not self._destructive_boundary_crossed
                and self._wechat_responsive
                and item_id not in self._non_retryable_item_ids
            ):
                source = next(
                    (
                        item
                        for item in (self._original_payload or {}).get("items", [])
                        if item.get("itemId") == params.get("itemId")
                    ),
                    None,
                )
                if source is not None:
                    self._retry_candidate = copy.deepcopy(source)
            if self._done > previous_done:
                self._completed_item_ids.add(str(params.get("itemId", "")))
            self.progressChanged.emit()
            self.currentStepChanged.emit()
            self.executionStateChanged.emit()
            if self._current_step in {"send_verified", "submit_verified"}:
                self._agent.call(
                    "recovery.approve",
                    {
                        "decision": "acknowledge",
                        "taskId": self._task_id,
                        "itemId": str(params.get("itemId", "")),
                    },
                )
        elif method == "task.finished" and params.get("taskId") == self._task_id:
            self._task_events.append(copy.deepcopy(params))
            self._finished_result = copy.deepcopy(params)
            cleanup = params.get("cleanup")
            self._cleanup_result = copy.deepcopy(cleanup) if isinstance(cleanup, dict) else {}
            for item in self._items._items:
                if item.result != "pending":
                    continue
                event = {
                    "itemId": item.item_id,
                    "outcome": "stopped",
                    "detail": "未执行：任务已结束，请返回编辑后重新选择",
                    "step": "",
                }
                self._items.apply_event(event)
                if self._kind == "friend_add":
                    self._friends.model.apply_event(event)
            self._pending_start_request_id = 0
            raw_done = int(params.get("done", 0))
            self._done = min(
                self._total,
                self._recovery_offset + raw_done,
            )
            self._set_active(False)
            self._set_phase(
                "error"
                if str(params.get("outcome", "")) == "error"
                else "done"
            )
            self.progressChanged.emit()
            self.executionStateChanged.emit()
        elif method == "agent.status":
            status = str(params.get("status", ""))
            if status == "recovered" and self._pending_resume:
                self._agent.inspect()
            elif status == "recovery_failed" and self._pending_resume:
                self._fail_recovery(
                    "微信恢复失败：" + str(params.get("detail", "未知错误"))
                )

    @Slot(result=bool)
    def retryFailedItem(self) -> bool:
        if not self.safeRetryAvailable or self._retry_candidate is None:
            return False
        payload = copy.deepcopy(self._retry_candidate)
        original = self._original_payload or {}
        return self._start(
            str(original.get("kind", self._kind)),
            [payload],
            copy.deepcopy(original.get("options", {})),
        )

    @Slot()
    def detectWechatRecovery(self) -> None:
        self._agent.inspect()

    @Slot(result=bool)
    def restartWechatAfterFailure(self) -> bool:
        if not self.wechatRestartAvailable:
            return False
        self._wechat_restart_attempted = True
        self.executionStateChanged.emit()
        try:
            self._agent.call(
                "recovery.approve",
                {
                    "decision": "restart_wechat",
                    "loginTimeout": self._settings.loginTimeout,
                },
            )
        except Exception as exc:
            self._wechat_restart_attempted = False
            self._set_error(f"请求重启微信失败：{exc}")
            self.executionStateChanged.emit()
            return False
        self._set_error("正在重启微信；完成登录后请检测自动化状态")
        return True

    @Slot()
    def _sync_wechat_health(self) -> None:
        if self._current_error_code != "WECHAT_UNRESPONSIVE" or self._active:
            return
        responsive = bool(self._agent.windowResponsive)
        if responsive == self._wechat_responsive:
            return
        self._wechat_responsive = responsive
        if responsive:
            self._current_detail = "微信窗口已恢复响应；请返回编辑后重新开始任务"
            self._set_error(self._current_detail)
        self.executionStateChanged.emit()

    @Slot(str, result=str)
    @Slot(str, str, result=str)
    def stepLabel(self, code: str, outcome: str = "") -> str:
        if outcome in {"error", "unknown"}:
            return FAILED_STEP_LABELS.get(code, "自动化步骤失败")
        return {**MESSAGE_STEP_LABELS, **FRIEND_STEP_LABELS}.get(code, code)

    @Slot(str, result=bool)
    def exportResults(self, file_url: str) -> bool:
        path = file_url.replace("file:///", "")
        if sys.platform == "win32":
            path = path.lstrip("/")
        try:
            os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
            with open(path, "w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["目标", "结果", "步骤", "详情", "耗时"])
                for item in self._items._items:
                    writer.writerow(
                        [
                            item.target,
                            item.result,
                            self.stepLabel(item.step_code, item.result),
                            item.detail,
                            item.duration,
                        ]
                    )
            return True
        except Exception as exc:
            self._set_error(str(exc))
            return False

    @Slot(str, result=bool)
    def exportDiagnostics(self, file_url: str) -> bool:
        path = file_url.replace("file:///", "")
        if sys.platform == "win32":
            path = path.lstrip("/")
        try:
            destination = Path(path)
            destination.parent.mkdir(parents=True, exist_ok=True)
            log_dir = default_log_dir()
            summary = {
                "schemaVersion": 1,
                "createdAt": datetime.now(timezone.utc).isoformat(),
                "platform": platform.platform(),
                "pythonVersion": platform.python_version(),
                "taskKind": self._kind,
                "taskPhase": self._phase,
                "taskCounts": {
                    "processed": self._done,
                    "total": self._total,
                    "success": self.successCount,
                    "failure": self.failureCount,
                    "unknown": self.unknownCount,
                },
                "agentConnected": self._agent.connected,
                "wechatVersion": self._agent.wechatVersion,
                "versionSupported": self._agent.versionSupported,
                "sessionReady": self._agent.sessionReady,
                "sessionGeneration": self._agent.sessionGeneration,
                "windowResponsive": self._agent.windowResponsive,
                "sessionErrorCode": self._current_error_code,
                "health": self._agent.healthSnapshot,
                "cleanup": self.cleanupResult,
            }
            with zipfile.ZipFile(
                destination, "w", compression=zipfile.ZIP_DEFLATED
            ) as archive:
                archive.writestr(
                    "diagnostic-summary.json",
                    json.dumps(_redact_diagnostic(summary), ensure_ascii=False, indent=2),
                )
                archive.writestr(
                    "task-events.jsonl",
                    "".join(json.dumps(_redact_diagnostic(event), ensure_ascii=False) + "\n" for event in self._task_events),
                )
                if log_dir.is_dir():
                    for candidate in sorted(log_dir.iterdir()):
                        if not candidate.is_file():
                            continue
                        if not (
                            candidate.name.startswith("uia-diagnostics.jsonl")
                            or candidate.name.startswith("agent-stderr.log")
                        ):
                            continue
                        archive.write(candidate, f"logs/{candidate.name}")
            return True
        except Exception as exc:
            self._set_error(f"导出诊断包失败：{exc}")
            return False


__all__ = [
    "AgentController",
    "FriendController",
    "FRIEND_STEP_LABELS",
    "MessageController",
    "MESSAGE_STEP_LABELS",
    "SettingsController",
    "TaskController",
]
