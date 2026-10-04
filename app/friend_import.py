from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

from .friend_templates import render_friend_content, split_name
from .constants import FRIEND_BATCH_LIMIT_DEFAULT, normalize_friend_batch_limit


ACCOUNT_HEADER = "账号"
NAME_HEADER = "姓名"
GREETING_HEADER = "打招呼语"
REMARK_HEADER = "备注"
PHONE_PATTERN = re.compile(r"^(?:\+?86)?1[3-9]\d{9}$")
WEIXIN_ID_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9_-]{5,31}$")


class FriendImportError(ValueError):
    pass


@dataclass
class FriendRecord:
    item_id: str
    account: str
    greeting: str = ""
    remark: str = ""
    valid: bool = True
    error: str = ""
    status: str = "pending"
    selected: bool = False
    name: str = ""
    # None = follow global; "" = explicitly no relationship.
    relationship: str | None = None
    relationship_source: str = "global"
    source_row: int = 0
    rendered_greeting: str | None = None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_account(value: str) -> str:
    account = "".join(value.strip().split())
    if PHONE_PATTERN.fullmatch(account):
        if account.startswith("+86"):
            account = account[3:]
        elif account.startswith("86") and len(account) == 13:
            account = account[2:]
        return account
    return account.casefold()


def account_error(value: str) -> str:
    account = "".join(value.strip().split())
    if not account:
        return "账号不能为空"
    if PHONE_PATTERN.fullmatch(account) or WEIXIN_ID_PATTERN.fullmatch(account):
        return ""
    return "账号必须是微信号或手机号，不能使用昵称"


def validate_records(records: Sequence[FriendRecord], default_greeting: str = "",
                     default_relationship: str = "妈妈") -> None:
    seen: set[str] = set()
    for row_number, record in enumerate(records, start=1):
        record.account = _cell_text(record.account)
        errors = []
        error = account_error(record.account)
        normalized = normalize_account(record.account)
        if not error and normalized in seen:
            error = "账号重复"
        if not error:
            seen.add(normalized)
        if error:
            errors.append(error)
        if not record.name.strip():
            errors.append("姓名不能为空（后缀前需有姓名）")
        try:
            record.rendered_greeting, record.remark = render_friend_content(
                record.name, record.relationship, record.greeting, default_greeting, default_relationship)
        except ValueError as exc:
            record.rendered_greeting = None
            _, record.remark = render_friend_content(record.name, record.relationship, "", "", default_relationship)
            errors.append(str(exc))
        record.valid = not errors
        record.error = f"第 {record.source_row or row_number} 行：" + "；".join(errors) if errors else ""
        if errors:
            record.selected = False


def _records_from_rows(rows: Iterable[Sequence[Any]], warnings: list[str] | None = None,
                       selection_limit: int = FRIEND_BATCH_LIMIT_DEFAULT) -> list[FriendRecord]:
    materialized = [list(row) for row in rows]
    if not materialized:
        raise FriendImportError("导入文件为空")
    headers = [_cell_text(value) for value in materialized[0]]
    missing = [name for name in (NAME_HEADER, ACCOUNT_HEADER) if name not in headers]
    if missing:
        raise FriendImportError("首行缺少必填列：" + "、".join(missing) + "；请下载新模板并补充")
    duplicates = sorted({name for name in headers if name and headers.count(name) > 1})
    if duplicates:
        raise FriendImportError("首行存在重复表头：" + "、".join(duplicates))
    if REMARK_HEADER in headers and warnings is not None:
        warnings.append("已忽略旧“备注”列；备注将由姓名和后缀自动生成")
    index = {name: headers.index(name) for name in headers if name}

    def value_at(row: Sequence[Any], header: str) -> str:
        column = index.get(header)
        if column is None or column >= len(row):
            return ""
        return _cell_text(row[column])

    records = []
    for row_number, row in enumerate(materialized[1:], start=2):
        values = [value_at(row, name) for name in headers]
        if not any(values):
            continue
        name, relationship = split_name(value_at(row, NAME_HEADER))
        records.append(
            FriendRecord(
                item_id=f"row-{row_number}",
                account=value_at(row, ACCOUNT_HEADER),
                greeting=value_at(row, GREETING_HEADER),
                name=name,
                relationship=relationship,
                relationship_source="auto" if relationship is not None else "global",
                source_row=row_number,
            )
        )
    validate_records(records)
    return records


def _csv_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FriendImportError("CSV 编码必须为 UTF-8 BOM 或 GB18030")
    try:
        return [list(row) for row in csv.reader(io.StringIO(text, newline=""), strict=True)]
    except csv.Error as exc:
        raise FriendImportError(f"CSV 格式损坏：{exc}") from exc


def _xlsx_rows(path: Path) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FriendImportError("缺少 openpyxl，无法读取 XLSX") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise FriendImportError(f"无法读取 XLSX：{exc}") from exc
    try:
        worksheet = workbook.worksheets[0]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def load_friend_records(source, *, warnings: list[str] | None = None,
                        selection_limit: int = FRIEND_BATCH_LIMIT_DEFAULT) -> list[FriendRecord]:
    if isinstance(source, (str, Path)):
        path = Path(source)
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            rows = _csv_rows(path)
        elif suffix == ".xlsx":
            rows = _xlsx_rows(path)
        else:
            raise FriendImportError("仅支持 CSV 或 XLSX 文件")
        return _records_from_rows(rows, warnings, selection_limit)
    return _records_from_rows(source, warnings, selection_limit)


__all__ = [
    "FriendImportError",
    "FriendRecord",
    "account_error",
    "load_friend_records",
    "normalize_account",
    "validate_records",
]
