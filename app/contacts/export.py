from __future__ import annotations

import csv
import io
import json
import os
import tempfile
from collections.abc import Callable
from concurrent.futures import CancelledError
from pathlib import Path
from threading import Event
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

from .data import CONTACT_FIELDS, CONTACT_HEADERS


_FORMATS = ("csv", "json", "xlsx")
_STEM = "\u5fae\u4fe1\u8054\u7cfb\u4eba"
_BATCH_SIZE = 128


class ExportCancelled(CancelledError):
    """The contact export was cancelled and its transaction rolled back."""


def _check_cancel(cancel: Event | None) -> None:
    if cancel is not None and cancel.is_set():
        raise ExportCancelled("Contact export cancelled")


def _temporary(destination: Path, suffix: str) -> Path:
    descriptor, name = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=suffix, dir=destination.parent)
    os.close(descriptor)
    return Path(name)


def _check_destination(destination: Path, overwrite: bool) -> None:
    if destination.is_dir():
        raise IsADirectoryError(str(destination))
    if os.path.lexists(destination) and not overwrite:
        raise FileExistsError(str(destination))


def _csv_text(value: str) -> str:
    if value.startswith(("\t", "\r", "\n")) or value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write_csv(path: Path, rows: list[dict[str, str]], advance: Callable[[], None]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(CONTACT_HEADERS)
        for row in rows:
            writer.writerow([_csv_text(row[field]) for field in CONTACT_FIELDS])
            advance()


def _write_json(path: Path, rows: list[dict[str, str]], advance: Callable[[], None]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        stream.write("[\n")
        for index, row in enumerate(rows):
            if index:
                stream.write(",\n")
            stream.write("  " + json.dumps(row, ensure_ascii=False))
            advance()
        stream.write("\n]\n")


class _CancelAwareFile:
    def __init__(self, stream, cancel):
        self._stream = stream
        self._cancel = cancel
        self._aborted = False

    def __getattr__(self, name):
        return getattr(self._stream, name)

    def write(self, value):
        if not self._aborted:
            try:
                _check_cancel(self._cancel)
            except ExportCancelled:
                # Permit ZipFile's close to finish its temporary archive safely.
                self._aborted = True
                raise
        return self._stream.write(value)


class _MemoryWorksheetWriter(WorksheetWriter):
    def __init__(self, sheet, stream, cancel, row_count):
        self._cancel = cancel
        self._row_count = row_count
        super().__init__(sheet, out=stream)

    def rows(self):
        rows = self.ws.iter_rows(min_row=1, max_row=self._row_count,
                                 min_col=1, max_col=len(CONTACT_FIELDS))
        for index, row in enumerate(rows, start=1):
            _check_cancel(self._cancel)
            yield index, row


class _MemoryExcelWriter(ExcelWriter):
    def __init__(self, workbook, archive, cancel, row_count):
        super().__init__(workbook, archive)
        self._cancel = cancel
        self._row_count = row_count

    def write_worksheet(self, sheet):
        # Both normal and write-only openpyxl defaults spool plaintext XML to
        # global temp. This export-specific writer keeps worksheet XML in RAM.
        with io.BytesIO() as stream:
            writer = _MemoryWorksheetWriter(sheet, stream, self._cancel, self._row_count)
            try:
                writer.write()
                sheet._rels = writer._rels
                stream.seek(0)
                with self._archive.open(sheet.path[1:], "w") as output:
                    while chunk := stream.read(65536):
                        _check_cancel(self._cancel)
                        output.write(chunk)
                self.manifest.append(sheet)
            finally:
                writer.close()


def _preserve_carriage_returns(path: Path, cancel: Event | None) -> None:
    repaired = _temporary(path, ".tmp")
    try:
        with ZipFile(path) as source, ZipFile(repaired, "w") as target:
            for entry in source.infolist():
                with source.open(entry) as reader, target.open(entry, "w") as writer:
                    while chunk := reader.read(65536):
                        _check_cancel(cancel)
                        # ElementTree emits literal CRs, which XML readers
                        # normalize to LF. Character references remain lossless.
                        if entry.filename.endswith(".xml"):
                            chunk = chunk.replace(b"\r", b"&#13;")
                        writer.write(chunk)
        _check_cancel(cancel)
        os.replace(repaired, path)
    finally:
        repaired.unlink(missing_ok=True)


def _write_xlsx(path: Path, rows: list[dict[str, str]], advance: Callable[[], None], cancel: Event | None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _STEM
    try:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:F{len(rows) + 1}"
        for column, width in zip("ABCDEF", (24, 28, 24, 34, 28, 48)):
            sheet.column_dimensions[column].width = width
        sheet.row_dimensions[1].height = 24
        sheet.append(CONTACT_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="26705A")
            cell.alignment = Alignment(vertical="center")
        for index, row in enumerate(rows, start=2):
            _check_cancel(cancel)
            for column, field in enumerate(CONTACT_FIELDS, start=1):
                if len(row[field]) > 32767:
                    raise ValueError(f"XLSX cell text exceeds the 32767 character length limit: {field}")
                cell = sheet.cell(index, column, row[field])
                cell.data_type = "s"
                cell.number_format = "@"
            advance()
        _check_cancel(cancel)
        with path.open("wb") as stream:
            with ZipFile(_CancelAwareFile(stream, cancel), "w", ZIP_DEFLATED, allowZip64=True) as archive:
                _MemoryExcelWriter(workbook, archive, cancel, len(rows) + 1).write_data()
        _check_cancel(cancel)
        if any("\r" in row[field] for row in rows for field in CONTACT_FIELDS):
            _preserve_carriage_returns(path, cancel)
    finally:
        workbook.close()


def export_contacts(
    records: list[dict],
    fmt: str,
    target: Path,
    *,
    overwrite: bool = False,
    cancel: Event | None = None,
    progress: Callable[[int, int], None] | None = None,
) -> list[str]:
    """Export frozen visible records synchronously, returning destination paths.

    progress(done, total) counts staged rows across formats, in bounded batches.
    Cancellation raises ExportCancelled (a concurrent.futures.CancelledError).
    JSON preserves raw text; CSV prefixes unsafe spreadsheet values; XLSX
    stores every value as a literal text cell. All writes finish before commit.
    """
    if fmt not in (*_FORMATS, "all"):
        raise ValueError(f"Unsupported contact export format: {fmt}")
    if not records:
        raise ValueError("No contacts to export")
    _check_cancel(cancel)

    rows = []
    for row in records:
        _check_cancel(cancel)
        rows.append({field: "" if row.get(field) is None else str(row[field]) for field in CONTACT_FIELDS})
    formats = _FORMATS if fmt == "all" else (fmt,)
    target = Path(target).absolute()
    destinations = [target / f"{_STEM}.{kind}" for kind in formats] if fmt == "all" else [target]
    for destination in destinations:
        _check_destination(destination, overwrite)

    total = len(rows) * len(formats)
    done = 0

    def report() -> None:
        _check_cancel(cancel)
        if progress is not None:
            progress(done, total)
        _check_cancel(cancel)

    def advance() -> None:
        nonlocal done
        _check_cancel(cancel)
        done += 1
        if done % _BATCH_SIZE == 0 or done == total:
            report()

    report()
    staged: dict[Path, Path] = {}
    backups: dict[Path, Path] = {}
    committed: list[Path] = []
    reserved: set[Path] = set()
    temporary_paths: set[Path] = set()
    retained_backups: set[Path] = set()
    writers = {"csv": _write_csv, "json": _write_json}
    try:
        for kind, destination in zip(formats, destinations):
            _check_cancel(cancel)
            destination.parent.mkdir(parents=True, exist_ok=True)
            stage = _temporary(destination, ".tmp")
            temporary_paths.add(stage)
            staged[destination] = stage
            if kind == "xlsx":
                _write_xlsx(stage, rows, advance, cancel)
            else:
                writers[kind](stage, rows, advance)
            _check_cancel(cancel)

        # Recheck the full set after staging, before the first original moves.
        for destination in destinations:
            _check_destination(destination, overwrite)
        for destination in destinations:
            _check_cancel(cancel)
            _check_destination(destination, overwrite)
            if not overwrite:
                descriptor = os.open(destination, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
                reserved.add(destination)
                os.close(descriptor)
            elif os.path.lexists(destination):
                backup = _temporary(destination, ".backup")
                temporary_paths.add(backup)
                os.replace(destination, backup)
                backups[destination] = backup
            os.replace(staged[destination], destination)
            committed.append(destination)
        _check_cancel(cancel)
    except BaseException as error:
        rollback_errors = []
        for destination in reversed(destinations):
            try:
                if destination in backups:
                    os.replace(backups[destination], destination)
                elif destination in committed or destination in reserved:
                    destination.unlink(missing_ok=True)
            except OSError as rollback_error:
                rollback_errors.append(rollback_error)
                if destination in backups:
                    retained_backups.add(backups[destination])
        for rollback_error in rollback_errors:
            error.add_note(f"Contact export rollback failed: {rollback_error}")
        for backup in retained_backups:
            error.add_note(f"Original contact export retained at: {backup}")
        raise
    finally:
        for path in temporary_paths - retained_backups:
            path.unlink(missing_ok=True)
    return [str(destination) for destination in destinations]
